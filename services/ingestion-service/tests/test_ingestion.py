import uuid
from pathlib import Path

from app.db.models.ingestion import IngestionChunk, IngestionRun
from app.services.ingestion_service import IngestionService


def test_clean_text_preserves_paragraphs():
    result = IngestionService.clean_text("  First   line  \n\n\n Second line \n\n Third")
    assert result == "First line\n\nSecond line\n\nThird"


def test_document_not_found(client):
    response = client.post(f"/ingestion/{uuid.uuid4()}")
    assert response.status_code == 404


def test_successful_ingestion_and_persistence(client, db_session, tmp_path):
    document_id = uuid.uuid4()
    path = Path(tmp_path) / f"{document_id}.docx"
    from io import BytesIO
    from docx import Document

    output = BytesIO()
    document = Document()
    document.add_paragraph("A document owned by Block 2.")
    document.save(output)
    path.write_bytes(output.getvalue())

    response = client.post(f"/ingestion/{document_id}")
    assert response.status_code == 200
    assert response.json()["status"] == "COMPLETED"
    assert response.json()["chunk_count"] == 1

    run = db_session.query(IngestionRun).one()
    chunk = db_session.query(IngestionChunk).one()
    assert run.document_id == str(document_id)
    assert run.status == "COMPLETED"
    assert chunk.chunk_index == 0
    assert chunk.content == "A document owned by Block 2."


def test_idempotent_ingestion_returns_existing_result(client, tmp_path):
    document_id = uuid.uuid4()
    (Path(tmp_path) / f"{document_id}.docx").write_bytes(b"not a docx")
    first = client.post(f"/ingestion/{document_id}")
    assert first.status_code == 500

    from io import BytesIO
    from docx import Document

    output = BytesIO()
    document = Document()
    document.add_paragraph("Retryable document")
    document.save(output)
    (Path(tmp_path) / f"{document_id}.docx").write_bytes(output.getvalue())
    second = client.post(f"/ingestion/{document_id}")
    assert second.status_code == 200
    assert second.json()["status"] == "COMPLETED"
    third = client.post(f"/ingestion/{document_id}")
    assert third.status_code == 200
    assert third.json()["message"] == "Document was already ingested."


def test_unsupported_document(client, tmp_path):
    document_id = uuid.uuid4()
    (Path(tmp_path) / f"{document_id}.txt").write_text("unsupported")
    response = client.post(f"/ingestion/{document_id}")
    assert response.status_code == 400
