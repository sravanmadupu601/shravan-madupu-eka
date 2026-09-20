from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.parser_service import DOCXParser, PDFParser, ParserFactory, UnsupportedDocumentError


def test_pdf_parser_extracts_text():
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    page.merge_page(page)
    writer.write(output)
    result = PDFParser().parse(output.getvalue(), "test.pdf", "application/pdf")
    assert result.text == ""


def test_docx_parser_extracts_paragraphs():
    output = BytesIO()
    document = DocxDocument()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(output)
    result = DOCXParser().parse(output.getvalue(), "test.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.text == "First paragraph\n\nSecond paragraph"


def test_parser_factory_rejects_unsupported_type():
    try:
        ParserFactory().select("test.txt", "text/plain")
        assert False
    except UnsupportedDocumentError as exc:
        assert "Unsupported document type" in str(exc)
