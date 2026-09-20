from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    page_text: list[tuple[int | None, str]]


class DocumentParser:
    def parse(self, file_content: bytes, filename: str, content_type: str) -> ParsedDocument:
        raise NotImplementedError


class PDFParser(DocumentParser):
    def parse(self, file_content: bytes, filename: str, content_type: str) -> ParsedDocument:
        reader = PdfReader(BytesIO(file_content))
        pages: list[tuple[int | None, str]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append((page_number, text))
        return ParsedDocument(
            text="\n\n".join(text for _, text in pages),
            page_text=pages,
        )


class DOCXParser(DocumentParser):
    def parse(self, file_content: bytes, filename: str, content_type: str) -> ParsedDocument:
        document = DocxDocument(BytesIO(file_content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        paragraphs = [paragraph for paragraph in paragraphs if paragraph]
        text = "\n\n".join(paragraphs)
        return ParsedDocument(text=text, page_text=[(None, text)] if text else [])


class ParserFactory:
    _parsers = {
        ".pdf": PDFParser(),
        ".docx": DOCXParser(),
    }

    def select(self, filename: str, content_type: str) -> DocumentParser:
        extension = Path(filename).suffix.lower()
        parser = self._parsers.get(extension)
        if parser is None:
            raise UnsupportedDocumentError(f"Unsupported document type: {extension or content_type}")
        return parser
